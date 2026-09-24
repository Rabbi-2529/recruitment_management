"""
Read an interview-call file (CSV, Excel or a BD Jobs Word applicant list) into plain rows,
guess which column is which, and turn the rows into InterviewCallCandidate records.

Nothing here touches the main Candidate table - a registered candidate with the same phone
number is only *linked* for reference.
"""
import csv
import io
import re
import zipfile
from xml.etree import ElementTree

from .utils import is_valid_phone, normalize_phone

SUPPORTED_EXTENSIONS = (".csv", ".xlsx", ".xlsm", ".docx", ".txt")

# the fields an admin can map a column to
TARGET_FIELDS = [
    ("candidate_name", "Candidate Name"),
    ("phone", "Phone Number"),
    ("email", "Email"),
    ("candidate_ref", "Candidate ID / Roll"),
    ("department_text", "Department"),
    ("education", "Education"),
    ("experience", "Experience"),
]
REQUIRED_FIELDS = {"phone"}

# header words that usually mean a field (compared without spaces / punctuation, lower case)
FIELD_ALIASES = {
    "candidate_name": ["candidatename", "name", "fullname", "applicantname", "studentname"],
    "phone": ["phone", "phonenumber", "mobile", "mobilenumber", "mobileno", "phoneno", "contact",
              "contactnumber", "contactno", "cell", "cellphone"],
    "email": ["email", "emailaddress", "mail", "emailid"],
    "candidate_ref": ["candidateid", "applicantid", "applicationid", "roll", "rollno", "rollnumber", "id",
                      "regno", "registrationno", "registrationnumber"],
    "department_text": ["department", "dept", "post", "position", "jobtitle"],
    "education": ["education", "qualification", "degree", "institute", "institution"],
    "experience": ["experience", "totalexperience", "careersummary", "workexperience"],
}

PHONE_RE = re.compile(r"(?<!\d)(?:\+?880|0)?1[3-9]\d{8}(?!\d)")
EMAIL_RE = re.compile(r"[A-Za-z0-9._%+-]+@[A-Za-z0-9.-]+\.[A-Za-z]{2,}")
W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"


class ImportError_(Exception):
    """A readable problem with the uploaded file."""


# ------------------------------------------------------------------ reading
def header_key(text):
    return re.sub(r"[^a-z0-9]", "", str(text or "").lower())


def read_rows(file_obj, filename):
    """Return (headers, rows) where rows are dicts keyed by header."""
    name = (filename or "").lower()
    if name.endswith(".csv"):
        table = _read_csv(file_obj)
    elif name.endswith((".xlsx", ".xlsm")):
        table = _read_xlsx(file_obj)
    elif name.endswith(".docx"):
        return _read_docx(file_obj)
    elif name.endswith(".txt"):
        return _read_txt(file_obj)
    else:
        raise ImportError_("Unsupported file. Upload a .xlsx, .csv, .txt or .docx (BD Jobs applicant list) file.")
    return _table_to_rows(table)


def _clean(value):
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)  # Excel stores phone numbers as numbers
    return str(value).strip()


def _table_to_rows(table):
    table = [[_clean(c) for c in row] for row in table if any(_clean(c) for c in row)]
    if not table:
        raise ImportError_("The file is empty.")
    headers, seen = [], {}
    for i, raw in enumerate(table[0]):
        label = raw or f"Column {i + 1}"
        if label in seen:  # two columns called "Phone" -> "Phone", "Phone (2)"
            seen[label] += 1
            label = f"{label} ({seen[label]})"
        else:
            seen[label] = 1
        headers.append(label)
    rows = []
    for raw in table[1:]:
        rows.append({h: (raw[i] if i < len(raw) else "") for i, h in enumerate(headers)})
    return headers, rows


def _decode(data):
    if isinstance(data, str):
        return data
    for encoding in ("utf-8-sig", "utf-16", "cp1252"):
        try:
            return data.decode(encoding)
        except UnicodeError:
            continue
    return data.decode("latin-1")


def _read_csv(file_obj):
    return list(csv.reader(io.StringIO(_decode(file_obj.read()))))


TXT_HEADERS = ["Name", "Phone", "Email", "Line"]


def _read_txt(file_obj):
    """
    A .txt file is either a table (comma / tab / semicolon / | separated, first line = headers)
    or a plain list with one person per line, e.g. "Rahim Ahmed - 01911111111 - rahim@mail.com".
    """
    text = _decode(file_obj.read())
    lines = [line for line in text.splitlines() if line.strip()]
    if not lines:
        raise ImportError_("The text file is empty.")
    try:
        dialect = csv.Sniffer().sniff("\n".join(lines[:20]), delimiters=",\t;|")
        table = list(csv.reader(io.StringIO("\n".join(lines)), dialect))
    except csv.Error:
        table = []
    # a real table: several columns on most lines and a header row without phone numbers in it
    if (table and len(table[0]) > 1
            and sum(1 for row in table if len(row) == len(table[0])) >= 0.8 * len(table)
            and not PHONE_RE.search(_squash_digits(" ".join(table[0])))):
        return _table_to_rows(table)

    rows = []
    for line in lines:
        squashed = _squash_digits(line)
        phone_match = PHONE_RE.search(squashed)
        email_match = EMAIL_RE.search(line)
        name = squashed
        for match in (phone_match, email_match):
            if match:
                name = name.replace(match.group(0), " ")
        name = re.sub(r"(?i)\b(name|phone|mobile|email|cell)\s*[:=]", " ", name)
        name = re.sub(r"^[\s\d.)#-]+(?=\D)", "", name)  # "1. Rahim" -> "Rahim"
        name = " ".join(re.sub(r"[,;|\t\-–]+", " ", name).split())
        rows.append({
            "Name": name[:200],
            "Phone": phone_match.group(0) if phone_match else "",
            "Email": email_match.group(0) if email_match else "",
            "Line": line.strip(),
        })
    return list(TXT_HEADERS), rows


def _read_xlsx(file_obj):
    from openpyxl import load_workbook

    try:
        wb = load_workbook(file_obj, read_only=True, data_only=True)
    except Exception as exc:  # not a real Excel file
        raise ImportError_(f"Could not read the Excel file: {exc}")
    ws = wb.active
    rows = [list(r) for r in ws.iter_rows(values_only=True)]
    wb.close()
    return rows


def _cell_text(tc):
    lines = []
    for p in tc.iter(W + "p"):
        parts = []
        for el in p.iter():
            if el.tag == W + "t":
                parts.append(el.text or "")
            elif el.tag in (W + "br", W + "cr"):
                parts.append("\n")
            elif el.tag == W + "tab":
                parts.append(" ")
        lines.append("".join(parts))
    return "\n".join(lines).strip()


def _read_docx(file_obj):
    try:
        with zipfile.ZipFile(file_obj) as z:
            root = ElementTree.fromstring(z.read("word/document.xml"))
    except (zipfile.BadZipFile, KeyError) as exc:
        raise ImportError_(f"Could not read the Word file: {exc}")
    tables = []
    for tbl in root.iter(W + "tbl"):
        rows = [[_cell_text(tc) for tc in tr.findall(W + "tc")] for tr in tbl.findall(W + "tr")]
        if rows:
            tables.append(rows)
    if not tables:
        raise ImportError_("No table was found in the Word file.")
    table = max(tables, key=len)  # the applicant list is the big table
    header = [header_key(h) for h in table[0]]
    if "name" in header and ("careersummary" in header or "experienceandapplicationstatus" in header):
        return _bdjobs_rows(table)
    flat = [[cell.replace("\n", ", ") for cell in row] for row in table]
    return _table_to_rows(flat)


BDJOBS_HEADERS = [
    "SL", "Name", "Age", "Location", "Education", "Job Matching", "Phone", "Email",
    "Career Summary", "Total Experience", "Expected Salary", "Total Applied", "Applied On", "Remarks",
]


def _bdjobs_rows(table):
    """BD Jobs "Applicant List Summary": the Name cell holds name, age, location, education, phone, email."""
    header = [header_key(h) for h in table[0]]
    col = {key: header.index(key) for key in header}

    def cell(row, key):
        i = col.get(key)
        return row[i] if i is not None and i < len(row) else ""

    rows = []
    for raw in table[1:]:
        lines = [line.strip().rstrip(",") for line in cell(raw, "name").split("\n") if line.strip()]
        if not lines:
            continue
        record = {h: "" for h in BDJOBS_HEADERS}
        record["SL"] = cell(raw, "sl")
        record["Name"] = lines[0]
        others = []
        for line in lines[1:]:
            low = line.lower()
            if low.startswith("age:"):
                record["Age"] = line.split(":", 1)[1].strip()
            elif low.startswith("job matching"):
                record["Job Matching"] = line.split(":", 1)[-1].strip()
            elif EMAIL_RE.search(line):
                record["Email"] = EMAIL_RE.search(line).group(0)
            elif PHONE_RE.search(_squash_digits(line)):
                record["Phone"] = line
            else:
                others.append(line)
        if others:
            record["Location"] = others[0]
            record["Education"] = " — ".join(others[1:])
        record["Career Summary"] = cell(raw, "careersummary").replace("\n", " | ")
        for line in cell(raw, "experienceandapplicationstatus").split("\n"):
            if ":" not in line:
                continue
            key, value = (x.strip() for x in line.split(":", 1))
            target = {"totalexperience": "Total Experience", "salary": "Expected Salary",
                      "totalapplied": "Total Applied"}.get(header_key(key))
            if target:
                record[target] = value
        record["Applied On"] = cell(raw, "appliedon")
        record["Remarks"] = cell(raw, "remarks")
        rows.append(record)
    if not rows:
        raise ImportError_("The BD Jobs table has no applicants.")
    return list(BDJOBS_HEADERS), rows


# ------------------------------------------------------------------ understanding the columns
def guess_mapping(headers):
    """{target field: header} using common column names; each header is used once."""
    mapping, used = {}, set()
    keys = {h: header_key(h) for h in headers}
    for field, aliases in FIELD_ALIASES.items():
        # exact column names first (in alias order), then a column name that contains a main alias
        match = next((h for alias in aliases for h in headers if keys[h] == alias and h not in used), None)
        if match is None:
            match = next((h for h in headers if h not in used and any(a in keys[h] for a in aliases[:3])), None)
        if match is not None:
            mapping[field] = match
            used.add(match)
    return mapping


def _squash_digits(text):
    """"01741-501 721" -> "01741501721" so numbers written with spaces or dashes are found."""
    return re.sub(r"(?<=\d)[\s\-]+(?=\d)", "", text)


def extract_phone(value):
    """The first valid Bangladeshi mobile number in a cell (cells can hold several), as 01XXXXXXXXX."""
    text = _squash_digits(_clean(value))
    for match in PHONE_RE.findall(text):
        phone = normalize_phone(match)
        if is_valid_phone(phone):
            return phone
    phone = normalize_phone(text)  # a plain number cell
    return phone if is_valid_phone(phone) else ""


def extract_email(value):
    match = EMAIL_RE.search(_clean(value))
    return match.group(0).lower() if match else ""


def build_people(call, headers, rows, mapping):
    """
    Create this call's InterviewCallCandidate rows.
    One phone per person; repeated phone numbers in the file are kept once.
    Returns {"created", "duplicates", "invalid"}.
    """
    from .models import Candidate, InterviewCallCandidate

    registered = {
        c.phone: c for c in Candidate.objects.filter(department=call.department).only("id", "phone")
    } if call.department_id else {}

    call.people.all().delete()
    seen, people = set(), []
    duplicates = invalid = 0
    mapped_headers = {h for h in mapping.values() if h}
    for number, row in enumerate(rows, start=2):  # row 1 is the header
        def value(field):
            header = mapping.get(field)
            return _clean(row.get(header, "")) if header else ""

        phone_raw = value("phone")
        phone = extract_phone(phone_raw)
        name = value("candidate_name")
        if not phone and not name:
            continue  # blank line
        if phone and phone in seen:
            duplicates += 1  # same person again - never send twice
            continue
        extra = {h: _clean(v) for h, v in row.items() if h not in mapped_headers and _clean(v)}
        for field in ("education", "experience"):
            if value(field):
                extra[dict(TARGET_FIELDS)[field]] = value(field)
        person = InterviewCallCandidate(
            interview_call=call,
            row_number=number,
            candidate_name=name[:200],
            phone=phone,
            phone_raw=phone_raw[:255],
            email=extract_email(value("email")) or value("email")[:254],
            candidate_ref=value("candidate_ref")[:100],
            department_text=value("department_text")[:200],
            additional_data=extra,
            matched_candidate=registered.get(phone),
        )
        if phone:
            seen.add(phone)
        else:
            person.is_valid = False
            person.sms_status = InterviewCallCandidate.SmsStatus.INVALID
            person.error = "No valid phone number" + (f" ({phone_raw})" if phone_raw else "")
            invalid += 1
        people.append(person)
    InterviewCallCandidate.objects.bulk_create(people)
    call.total_rows = len(rows)
    call.duplicate_rows = duplicates
    call.save(update_fields=["total_rows", "duplicate_rows"])
    return {"created": len(people), "duplicates": duplicates, "invalid": invalid}
