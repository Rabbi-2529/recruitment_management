import csv
import io
import re

HEADER_ALIASES = {
    "name": {"name", "candidate", "candidatename", "fullname", "applicantname"},
    "phone": {"phone", "phonenumber", "mobile", "mobilenumber", "mobileno", "phoneno", "contact", "contactnumber"},
    "department": {"department", "dept", "departmentname", "post", "position"},
    "status": {"status", "result"},
    "email": {"email", "emailaddress", "mail"},
    "written_mark": {"written", "writtenmark", "writtenexam"},
    "rating": {"rating", "stars"},
    "expected_salary": {"expectedsalary", "salary", "expectedsalarybdt"},
    "available_date": {"availabledate", "available", "availablefrom", "joiningdate"},
    "note": {"note", "notes", "remarks", "remark", "comment", "comments"},
}

STATUS_ALIASES = {
    "pending": "pending",
    "": "pending",
    "selected": "selected",
    "select": "selected",
    "yes": "selected",
    "waiting": "waiting",
    "wait": "waiting",
    "waitlist": "waiting",
    "waitinglist": "waiting",
    "hold": "waiting",
    "onhold": "waiting",
    "notselected": "not_selected",
    "rejected": "not_selected",
    "no": "not_selected",
}


def normalize_phone(value):
    """Normalize Bangladeshi phone numbers to 11-digit local form (01XXXXXXXXX)."""
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    digits = re.sub(r"\D", "", str(value))
    if digits.startswith("880") and len(digits) == 13:
        digits = "0" + digits[3:]
    elif len(digits) == 10 and digits.startswith("1"):
        # Excel drops the leading zero of numeric cells
        digits = "0" + digits
    return digits


def is_valid_phone(phone):
    return bool(re.fullmatch(r"01[3-9]\d{8}", phone))


BD_LOCAL_RE = re.compile(r"01[3-9]\d{8}")  # 01XXXXXXXXX   (11 digits)
BD_INTL_RE = re.compile(r"8801[3-9]\d{8}")  # 8801XXXXXXXXX (13 digits)
PHONE_ERROR = "Invalid phone number. Use 01XXXXXXXXX (11 digits) or 8801XXXXXXXXX (13 digits)."


def validate_bd_phone(raw):
    """
    Strict check for a typed phone number. Accepts only:
      01XXXXXXXXX    - 11 digits, operator 013-019
      8801XXXXXXXXX  - 13 digits, operator 013-019
    Returns the number in 01XXXXXXXXX form, or None if invalid.
    """
    value = re.sub(r"[\s-]", "", str(raw or ""))
    if BD_LOCAL_RE.fullmatch(value):
        return value
    if BD_INTL_RE.fullmatch(value):
        return value[2:]
    return None


def _header_key(raw):
    key = re.sub(r"[^a-z]", "", str(raw or "").lower())
    for field, aliases in HEADER_ALIASES.items():
        if key in aliases:
            return field
    return None


def parse_mark(value, maximum=5):
    """Return a Decimal mark between 0 and maximum (max 2 decimals), or None if invalid."""
    from decimal import Decimal, InvalidOperation

    if isinstance(value, float):
        value = repr(value)
    try:
        mark = Decimal(str(value).strip())
    except (InvalidOperation, ValueError):
        return None
    if not mark.is_finite() or not (0 <= mark <= Decimal(maximum)):
        return None
    if mark != mark.quantize(Decimal("0.01")):
        return None
    return mark.quantize(Decimal("0.01"))


def parse_whole_mark(value, maximum=5):
    """
    Interview marks are whole numbers only: 0, 1, 2, 3, 4, 5.
    "4" (or 4.0 from an Excel cell) -> 4.  "4.5", "abc", 6 -> None (rejected).
    """
    from decimal import Decimal, InvalidOperation

    if isinstance(value, bool) or value is None:
        return None
    try:
        number = Decimal(str(value).strip())
    except (InvalidOperation, ValueError):
        return None
    if not number.is_finite() or number != number.to_integral_value():
        return None
    number = int(number)
    return number if 0 <= number <= int(maximum) else None


def parse_rating(value):
    try:
        rating = int(float(str(value).strip()))
    except ValueError:
        return None
    return rating if 1 <= rating <= 5 and float(str(value).strip()) == rating else None


def parse_salary(value):
    if isinstance(value, (int, float)):
        return int(value) if value >= 0 and float(value).is_integer() else None
    raw = re.sub(r"[,\s৳]|tk|bdt", "", str(value).strip().lower())
    return int(raw) if raw.isdigit() else None


def parse_date(value):
    import datetime

    if isinstance(value, datetime.datetime):
        return value.date()
    if isinstance(value, datetime.date):
        return value
    raw = str(value).strip()
    for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y", "%d.%m.%Y", "%d %b %Y", "%d %B %Y"):
        try:
            return datetime.datetime.strptime(raw, fmt).date()
        except ValueError:
            continue
    return None


def parse_status(value):
    key = re.sub(r"[^a-z]", "", str(value or "").lower())
    return STATUS_ALIASES.get(key)


def read_candidate_file(uploaded_file):
    """
    Read a CSV or XLSX file and return (rows, error).
    Each row is a dict with keys name/phone/department/status and a 'line' number.
    """
    filename = uploaded_file.name.lower()
    try:
        if filename.endswith(".csv"):
            raw_rows = _read_csv(uploaded_file)
        elif filename.endswith((".xlsx", ".xlsm")):
            raw_rows = _read_xlsx(uploaded_file)
        else:
            return [], "Unsupported file type. Please upload a .csv or .xlsx file."
    except Exception as exc:  # malformed file
        return [], f"Could not read the file: {exc}"

    if not raw_rows:
        return [], "The file is empty."

    header = [_header_key(h) for h in raw_rows[0]]
    missing = {"name", "phone", "department"} - set(header)
    if missing:
        return [], (
            "Missing required column(s): " + ", ".join(sorted(missing))
            + ". The first row must contain: Name, Phone, Department (Status is optional)."
        )

    rows = []
    for line_no, raw in enumerate(raw_rows[1:], start=2):
        record = {"line": line_no}
        for idx, field in enumerate(header):
            if field and idx < len(raw):
                record[field] = raw[idx]
        if not any(str(record.get(f) or "").strip() for f in ("name", "phone", "department")):
            continue  # blank line
        rows.append(record)
    return rows, None


def _read_csv(uploaded_file):
    data = uploaded_file.read()
    for encoding in ("utf-8-sig", "cp1252"):
        try:
            text = data.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    return [row for row in csv.reader(io.StringIO(text))]


def _read_xlsx(uploaded_file):
    from openpyxl import load_workbook

    wb = load_workbook(uploaded_file, read_only=True, data_only=True)
    ws = wb.active
    rows = [list(r) for r in ws.iter_rows(values_only=True)]
    wb.close()
    return rows
